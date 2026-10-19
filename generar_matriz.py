"""Genera la matriz de trazabilidad (xlsx) a partir del archivo de historias.

Uso:
    python generar_matriz.py [historias.md] [matriz_v1.xlsx] [salida.xlsx]
"""
import re
import sys
from collections import OrderedDict

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from create_issues import parse_readme

MD_PATH = sys.argv[1] if len(sys.argv) > 1 else "issues_GeDoC.md"
V1_PATH = sys.argv[2] if len(sys.argv) > 2 else "../../docs/trazabilidad_ISO7101_Garrahan.xlsx"
OUT_PATH = sys.argv[3] if len(sys.argv) > 3 else "../../docs/trazabilidad_ISO7101_Garrahan_v2.xlsx"

HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(bold=True, color="FFFFFF")
WRAP = Alignment(wrap_text=True, vertical="top")

TRAZ_RE = re.compile(r"\*\*Trazabilidad\*\*: (.*?) \| Story points: (\d+) \| Matriz v1: (.*)")
CA_RE = re.compile(
    r"- CA (\d+): (.*)\n\s+- (Dado.*)\n\s+- (Cuando.*)\n\s+- (Entonces.*)"
)


def parse_story(issue):
    body = issue["body"]
    hu_id, nombre = issue["title"].split(" ", 1)
    traz = TRAZ_RE.search(body)
    if not traz:
        raise ValueError(f"{hu_id}: falta la línea **Trazabilidad**")
    desc = re.search(r"\*\*Descripción\*\*: (.*)", body).group(1)
    req = re.search(r"\*\*Requisito\*\*: (.*)", body)
    funcionalidad = re.search(r"quiero \*\*(.*?)\*\*", desc).group(1)
    v1 = traz.group(3).strip()
    v1_num = int(re.search(r"\d+", v1).group()) if v1.startswith("N°") else None
    labels = issue["labels"]
    return {
        "id": hu_id,
        "nombre": nombre,
        "etapa": labels[0].capitalize() if labels[0].startswith("etapa") else "",
        "epica": next((l.split(":", 1)[1] for l in labels if l.startswith("epica:")), ""),
        "clausula": traz.group(1).strip(),
        "sp": int(traz.group(2)),
        "v1": v1,
        "v1_num": v1_num,
        "requisito": req.group(1) if req else "",
        "funcionalidad": funcionalidad[0].upper() + funcionalidad[1:],
        "historia": desc.replace("**", ""),
        "friccion": "PUNTO DE FRICCIÓN" in traz.group(1),
        "cas": [
            {"n": int(m[0]), "titulo": m[1], "dado": m[2], "cuando": m[3], "entonces": m[4]}
            for m in CA_RE.findall(body)
        ],
    }


def leer_v1(path):
    """Devuelve {N°: funcionalidad} y el texto de la hoja Metodología de la v1."""
    wb = load_workbook(path)
    funcionalidades = {}
    for row in wb["Trazabilidad"].iter_rows(min_row=2, values_only=True):
        if isinstance(row[0], int):
            funcionalidades[row[0]] = row[3]
    metodologia = [row[0] for row in wb["Metodología"].iter_rows(values_only=True)]
    return funcionalidades, metodologia


def escribir_hoja(ws, headers, rows, widths):
    ws.append(headers)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = WRAP
    for row in rows:
        ws.append(row)
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = WRAP
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def hoja_metodologia(ws, metodologia_v1, stories):
    ws.column_dimensions["A"].width = 140
    total_ca = sum(len(s["cas"]) for s in stories)
    nuevas = sum(1 for s in stories if s["v1"] == "Nueva")
    agregado = [
        None,
        "Cambios en la versión 2",
        f"- {len(stories)} historias de usuario ({nuevas} nuevas respecto de la v1) y {total_ca} casos de aceptación (antes uno por historia).",
        "- Cada historia tiene un identificador estable HU-xx y cada caso de aceptación HU-xx-CAn. La columna \"Matriz v1\" indica el N° de la historia original o \"Nueva\".",
        "- La historia N° 14 de la v1 (13 puntos) se dividió en HU-29 (planificación) y HU-30 (ejecución). La N° 18 (pista de auditoría) pasó de 5 a 8 puntos por el hallazgo del gap analysis de Alfresco. La N° 3 se generalizó a acuse de lectura de cualquier documento.",
        "- Se agregaron historias que la propuesta da por centrales y no estaban en la v1: ciclo de vida con estados, retiro automático de la versión obsoleta, búsqueda en el punto de uso, autenticación y gestión de usuarios, obsolescencia y retención, revisión periódica, plantillas, importación, período de lectura y reconstrucción a una fecha.",
        "- La fuente de verdad es repos/create-issues/issues_GeDoC.md; esta planilla se regenera con generar_matriz.py y los issues de GitHub se crean con create_issues.py.",
        None,
        "Etapas",
        "Etapa 1 = documentación mandatoria del Sistema de Gestión de la Calidad (meses 3 a 6 del plan: organigrama, ciclo de vida, doble aprobación, versionado, permisos, pista de auditoría) más las decisiones bloqueantes con Organización y Métodos. Etapa 2 = documentación operativa a escala y módulos de gestión de la calidad (meses 6 a 9).",
        None,
        "Roles",
        "Integrante del personal · Autor de un proceso · Responsable del proceso (aprobación técnica) · Responsable de Organización y Métodos (aprobación institucional, plantillas) · Responsable de calidad · Jefe de servicio · Auditor interno · Director del hospital · Responsable de Recursos Humanos · Administrador del sistema · Equipo de proyecto.",
    ]
    for value in metodologia_v1 + agregado:
        ws.append([value])
    for row in ws.iter_rows():
        cell = row[0]
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        if cell.value and len(cell.value) < 70 and not cell.value.startswith("-"):
            cell.font = Font(bold=True, size=14 if cell.row == 1 else 11)


def main():
    stories = [parse_story(i) for i in parse_readme(MD_PATH)]
    funcionalidades_v1, metodologia_v1 = leer_v1(V1_PATH)
    for s in stories:
        if s["v1_num"] in funcionalidades_v1:
            s["funcionalidad"] = funcionalidades_v1[s["v1_num"]]

    wb = Workbook()
    hoja_metodologia(wb.active, metodologia_v1, stories)
    wb.active.title = "Metodología"

    ws = wb.create_sheet("Trazabilidad")
    escribir_hoja(
        ws,
        ["ID", "Épica", "Etapa", "Cláusula ISO 7101", "Requisito de la norma (resumen)",
         "Funcionalidad del sistema", "Historia de usuario", "Story Points",
         "Cant. CA", "Matriz v1", "Issue GitHub"],
        [[s["id"], s["epica"].capitalize(), s["etapa"], s["clausula"], s["requisito"],
          s["funcionalidad"], s["historia"], s["sp"], len(s["cas"]), s["v1"], None]
         for s in stories],
        [8, 18, 9, 22, 45, 45, 55, 8, 8, 10, 10],
    )
    last = ws.max_row
    ws.append([None] * 6 + ["TOTAL", f"=SUM(H2:H{last})", f"=SUM(I2:I{last})"])
    for col in ("G", "H", "I"):
        ws[f"{col}{last + 1}"].font = Font(bold=True)
    ws.auto_filter.ref = f"A1:K{last}"

    ws = wb.create_sheet("Casos de aceptación")
    escribir_hoja(
        ws,
        ["ID CA", "ID HU", "Historia", "Caso", "Dado", "Cuando", "Entonces", "Estado de prueba"],
        [[f"{s['id']}-CA{c['n']}", s["id"], s["nombre"], c["titulo"],
          c["dado"], c["cuando"], c["entonces"], "Pendiente"]
         for s in stories for c in s["cas"]],
        [13, 8, 30, 30, 45, 45, 55, 14],
    )

    cobertura = OrderedDict()
    for s in stories:
        # "7.5.2 b" y "7.5.5.i" se normalizan a "7.5.2.b" / "7.5.5.i"
        clausulas = [re.sub(r"\s+", ".", c) for c in
                     re.findall(r"\d+(?:\.\d+)+(?:[.\s][a-z]\b)?", s["clausula"])]
        if "Disposición 186" in s["clausula"]:
            clausulas.append("Disposición 186/DME/2022")
        for cl in clausulas or [s["clausula"]]:
            cobertura.setdefault(cl, []).append(s)

    def clave(cl):
        if not cl[0].isdigit():
            return [999, cl]
        return [int(p) if p.isdigit() else ord(p) / 1000 for p in cl.split(".")]

    ws = wb.create_sheet("Cobertura por cláusula")
    escribir_hoja(
        ws,
        ["Cláusula", "Historias que la cubren", "Cant. HU", "Cant. CA", "Story Points"],
        [[cl, ", ".join(s["id"] for s in ss), len(ss),
          sum(len(s["cas"]) for s in ss), sum(s["sp"] for s in ss)]
         for cl, ss in sorted(cobertura.items(), key=lambda kv: clave(kv[0]))],
        [32, 50, 10, 10, 12],
    )

    ws = wb.create_sheet("Puntos de fricción")
    escribir_hoja(
        ws,
        ["ID", "Decisión pendiente", "Tensión", "Responsables", "Estado", "Fecha de decisión", "Constancia"],
        [[s["id"], s["nombre"], s["requisito"],
          "Organización y Métodos + Gerencia de Sistemas" if "SIG" in s["nombre"] else "Organización y Métodos",
          "Abierta", None, None]
         for s in stories if s["friccion"]],
        [8, 40, 70, 30, 10, 14, 30],
    )

    wb.save(OUT_PATH)
    print(f"{len(stories)} historias, {sum(len(s['cas']) for s in stories)} CA, "
          f"{sum(s['sp'] for s in stories)} SP -> {OUT_PATH}")


if __name__ == "__main__":
    main()
